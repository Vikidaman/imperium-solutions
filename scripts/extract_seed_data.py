import json
import openpyxl
import datetime

wb = openpyxl.load_workbook('Interstellar 2.0 Mastersheet.xlsx', read_only=True, data_only=True)

def clean_val(v, default=0.0):
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        if v.startswith('#') or v.strip() == '-' or v.strip() == '':
            return default
        try:
            return float(v)
        except ValueError:
            return default
    return default

# 1. Extract Time Manager
ws_tm = wb['Time Manager']
tm_rows = list(ws_tm.iter_rows(min_row=1, max_row=10, max_col=10, values_only=True))
time_manager = {
    "intervalDaysPerHour": clean_val(tm_rows[2][2], 90.0),
    "daysElapsed": clean_val(tm_rows[1][5], 0.0),
    "quarterNumber": int(clean_val(tm_rows[2][5], 1)),
    "crisisStartDate": "2025-01-01",
    "crisisDate": "2025-01-01",
    "nextPaymentDate": "2025-04-01",
    "isLive": False
}

# 2. Extract Political Parties
ws_pp = wb['Political Parties']
pp_rows = list(ws_pp.iter_rows(min_row=2, max_row=20, max_col=16, values_only=True))
political_parties = []
for r in pp_rows:
    if r and r[2] is not None:
        political_parties.append({
            "nation": str(r[1] or "Global"),
            "name": str(r[2]),
            "currentSupport": clean_val(r[3], 0.25),
            "incomeTaxPreference": clean_val(r[4], 0.2),
            "salesTaxPreference": clean_val(r[5], 0.1),
            "corporateTaxRate": clean_val(r[6], 0.2),
            "subsidyPreference": str(r[7] or "None"),
            "nationalizationPreference": str(r[8] or "No Nationalization"),
            "workerShareOfGdp": clean_val(r[9], 0.45),
            "budgetAsPercentOfGdp": clean_val(r[10], 0.15),
            "federalisedPolicy": str(r[11] or "Confederated"),
            "localCrimePolicy": str(r[12] or "Beat Cops"),
            "healthcarePolicy": str(r[13] or "Private Healthcare"),
            "politicalAlignment": str(r[14] or "Centre")
        })

# 3. Extract Companies Companion (Sectors & Country K-Factor Scalers)
ws_cc = wb['Companies Companion']
cc_rows = list(ws_cc.iter_rows(min_row=2, max_row=15, max_col=12, values_only=True))
sectors_config = []
country_scalers = {}
for r in cc_rows:
    if r[1] is not None:
        sectors_config.append({
            "name": str(r[1]),
            "sentiment": clean_val(r[2], 50.0),
            "kFactorBonus": clean_val(r[5], 0.3),
            "cbpf": clean_val(r[6], 0.0)
        })
    if len(r) > 9 and r[9] is not None:
        country_scalers[str(r[9])] = clean_val(r[10], 1500.0)

# 4. Extract Production Methods Companion (Tech Eras & 202 Goods)
ws_pmc = wb['Production Methods Companion']
pmc_rows = list(ws_pmc.iter_rows(min_row=1, max_row=69, max_col=210, values_only=True))

tech_eras = []
for r_idx in range(4, 13):
    row = pmc_rows[r_idx]
    if row[1] is not None:
        tech_eras.append({
            "name": str(row[1]),
            "unlocked": bool(row[2])
        })

goods_catalog = []
header_row = pmc_rows[0]
summary_row = pmc_rows[1]
sector_row = pmc_rows[2]
good_type_row = pmc_rows[3]
era_row = pmc_rows[4]

seen_good_names = set()
for col_idx in range(4, len(header_row)):
    g_name = header_row[col_idx]
    if g_name is not None:
        raw_name = str(g_name).strip()
        g_type = str(good_type_row[col_idx] or "General").strip()
        if raw_name in seen_good_names:
            raw_name = f"{raw_name} ({g_type})"
        seen_good_names.add(raw_name)
        inputs_raw = str(summary_row[col_idx] or "")
        inputs = [x.strip() for x in inputs_raw.split(",") if x.strip()]
        goods_catalog.append({
            "id": col_idx - 3,
            "name": raw_name,
            "sector": str(sector_row[col_idx] or "Industry & Manufacturing GDP").strip(),
            "goodType": g_type,
            "era": str(era_row[col_idx] or "Wood Tools Era").strip(),
            "inputs": inputs
        })

# 5. Extract Nations from Earth Full Data & Mars Full Data
nation_sheet_map = {
    "USA": ("USA Companies", "USA Market", "Earth"),
    "China": ("China Companies", "China Market", "Earth"),
    "Japan": ("Japan Companies", "Japan Market", "Earth"),
    "European Union": ("EU Companies", "EU Market", "Earth"),
    "India": ("India Companies", "India Market", "Earth"),
    "Russia": ("Russia Companies", "Russia Market", "Earth"),
    "Botswana": ("Botswana Companies", "Botswana Market", "Earth"),
    "Brazil": ("Brazil Companies", "Brazil Market", "Earth"),
    "Australia": ("Australia Companies", "Australia Market", "Earth"),
    "Mexico": ("Mexico Companies", "Mexico Market", "Earth"),
    "Olympus": ("Olympus Companies", "Olympus Market", "Mars"),
    "Heinan": ("Heinan Companies", "Heinan Market", "Mars"),
    "Mitsubishi": ("Mitsubishi Companies", "Mitsubishi Market", "Mars"),
    "Sumitomo": ("Sumitomo Companies", "Sumitomo Market", "Mars"),
    "I.F.C": ("IFC Companies", "IFC Market", "Mars"),
}

nations = []
for planet_sheet, planet_name in [("Earth Full Data", "Earth"), ("Mars Full Data", "Mars")]:
    ws_p = wb[planet_sheet]
    p_rows = list(ws_p.iter_rows(min_row=3, max_row=15, max_col=40, values_only=True))
    for r in p_rows:
        if r and r[1] is not None:
            n_name = str(r[1]).strip()
            if n_name not in nation_sheet_map:
                continue
            comp_sheet_name, mkt_sheet_name, _ = nation_sheet_map[n_name]
            
            # Read companies sheet metadata & companies
            ws_c = wb[comp_sheet_name]
            c_rows = list(ws_c.iter_rows(min_row=1, max_row=505, max_col=18, values_only=True))
            meta_row = c_rows[0]
            stock_exchange_name = str(meta_row[3] or f"{n_name} Stock Exchange")
            consumer_sentiment = clean_val(meta_row[7], 50.0)
            
            companies = []
            for idx, cr in enumerate(c_rows[2:], 1):
                if cr and cr[1] is not None:
                    companies.append({
                        "id": f"{n_name.lower().replace(' ', '_').replace('.', '')}_comp_{idx}",
                        "name": str(cr[1]).strip(),
                        "hqCountry": str(cr[2] or n_name).strip(),
                        "flagshipGood": str(cr[3] or "Tools").strip(),
                        "gdpSector": str(cr[4] or "Industry & Manufacturing GDP").strip(),
                        "hypeFactor": clean_val(cr[5], 50.0),
                        "valuation": clean_val(cr[6], 0.0),
                        "vgp": clean_val(cr[7], 0.0),
                        "employedPops": clean_val(cr[8], 10000.0),
                        "projectedRealGrowth": clean_val(cr[9], 0.0),
                        "govtOwnedShares": clean_val(cr[10], 0.0),
                        "numberOfShares": clean_val(cr[11], 50000000.0) if clean_val(cr[11], 0.0) > 0 else 3.0,
                        "dividendEfficiency": clean_val(cr[12], 0.6),
                        "dividendsPerShare": clean_val(cr[13], 0.0),
                        "productionMethod": str(cr[14] or "Tools").strip(),
                        "subsidised": bool(cr[15]) if cr[15] is not None else False,
                        "targetGrowth": clean_val(cr[16], 0.02),
                        "illegal": bool(cr[17]) if len(cr) > 17 and cr[17] is not None else False
                    })

            # Read market sheet
            ws_m = wb[mkt_sheet_name]
            m_rows = list(ws_m.iter_rows(min_row=2, max_row=205, max_col=25, values_only=True))
            market_items = []
            for mr in m_rows:
                if mr and mr[1] is not None:
                    market_items.append({
                        "goodName": str(mr[1]).strip(),
                        "hype": clean_val(mr[3], 50.0),
                        "ddMultiplier": clean_val(mr[9], 1.0),
                        "exportTariffRate": clean_val(mr[12], 0.0),
                        "importTariffRate": clean_val(mr[19], 0.0),
                        "illegal": bool(mr[23]) if len(mr) > 23 and mr[23] is not None else False
                    })

            nations.append({
                "id": n_name.lower().replace(" ", "_").replace(".", ""),
                "name": n_name,
                "planet": planet_name,
                "stockExchangeName": stock_exchange_name,
                "consumerSentiment": consumer_sentiment,
                "kFactorScaler": country_scalers.get(n_name, 1500.0),
                "workingPops": clean_val(r[6], 10000000.0),
                "literacyRate": clean_val(r[11], 0.50),
                "urbanisationIndex": clean_val(r[13], 0.50),
                "infrastructureEfficiency": clean_val(r[18], 0.75),
                "rulingParty": str(r[19] or "Republican").strip(),
                "federalCooperativity": clean_val(r[20], 0.0),
                "diplomaticExpenses": clean_val(r[33], 0.0),
                "treasuryBalance": clean_val(r[35], 0.0),
                "specializedIndustry": str(r[36] or "Industry & Manufacturing GDP").strip(),
                "crimePolicy": str(r[37] or "Profiling").strip(),
                "healthcarePolicy": str(r[38] or "Private Healthcare").strip(),
                "highImmigration": n_name in ["USA", "Olympus"],
                "companies": companies,
                "market": market_items
            })

# 6. Extract Stock Exchange Accounts
ws_se = wb['Stock Exchange Accounts']
se_rows = list(ws_se.iter_rows(min_row=1, max_row=25, max_col=35, values_only=True))
accounts = []
for acc_idx, start_col in enumerate([1, 6, 11, 16, 21, 26], 1):
    holder_name = str(se_rows[0][start_col] or f"Account Holder {acc_idx}")
    if holder_name == "Account Holder 1" and acc_idx > 1:
        holder_name = f"Account Holder {acc_idx}"
    div_pot = clean_val(se_rows[1][start_col + 1], 0.0)
    holdings = []
    for r in se_rows[4:20]:
        c_name = r[start_col]
        hq = r[start_col + 1]
        units = clean_val(r[start_col + 3], 0.0)
        if c_name and str(c_name) != "0" and units > 0:
            holdings.append({
                "companyName": str(c_name).strip(),
                "hqCountry": str(hq or "USA").strip(),
                "units": units
            })
    accounts.append({
        "id": f"acc_{acc_idx}",
        "holderName": holder_name,
        "dividendPot": div_pot,
        "cashBalance": 10000000000.0,
        "holdings": holdings
    })

seed_data = {
    "timeManager": time_manager,
    "politicalParties": political_parties,
    "sectorsConfig": sectors_config,
    "techEras": tech_eras,
    "goodsCatalog": goods_catalog,
    "nations": nations,
    "stockAccounts": accounts
}

with open('scripts/interstellar_seed.json', 'w') as f:
    json.dump(seed_data, f, indent=2)

print(f"Successfully extracted {len(nations)} nations, {len(goods_catalog)} goods, {sum(len(n['companies']) for n in nations)} total companies.")
