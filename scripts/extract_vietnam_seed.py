import json
import openpyxl

wb = openpyxl.load_workbook('Imperium Solutions Mastersheet.xlsx', read_only=True, data_only=True)
with open('scripts/interstellar_seed.json') as f:
    base_seed = json.load(f)

def clean_val(v, default=0.0):
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            return default
    return default

state_names = [
    "NorthEast Vietnam", "NorthWest Vietnam", "Red River Delta",
    "North Central Coast", "South Central Coast", "SouthEast Vietnam",
    "Mekong Delta", "Central Highlands"
]

ws_c = wb['Companies']
c_rows = list(ws_c.iter_rows(min_row=2, max_row=500, max_col=19, values_only=True))

companies_by_state = {s: [] for s in state_names}
for idx, r in enumerate(c_rows, 1):
    if r and len(r) > 8 and r[1] is not None and r[2] in companies_by_state:
        st = r[2]
        good = str(r[3] or "Rice").strip()
        comp_type = str(r[4] or "State-Owned Enterprise").strip()
        sector = str(r[5] or "Agricultural Sector GDP").strip()
        pe = clean_val(r[6], 1.0)
        val = clean_val(r[7], 1000000.0)
        vgp = clean_val(r[8], val)
        emp = clean_val(r[9], 10000.0) if len(r) > 9 else 10000.0
        growth = clean_val(r[10], 0.02) if len(r) > 10 else 0.02
        govt_shares = clean_val(r[11], 1.0) if len(r) > 11 else 1.0
        tot_shares = clean_val(r[13], 1.0) if len(r) > 13 and clean_val(r[13], 0.0) > 0 else 1.0
        div_eff = clean_val(r[14], 0.6) if len(r) > 14 else 0.6
        div_ps = clean_val(r[16], 0.0) if len(r) > 16 else 0.0
        prod_method = str(r[17] if len(r) > 17 and r[17] else "Tools").strip()
        companies_by_state[st].append({
            "id": f"vn_{st.lower().replace(' ', '_')}_{idx}",
            "name": f"{r[1]} ({comp_type})",
            "hqCountry": st,
            "flagshipGood": good if any(g['name'] == good for g in base_seed['goodsCatalog']) else "Rice",
            "gdpSector": sector,
            "hypeFactor": max(10.0, min(100.0, pe * 50.0)),
            "valuation": val,
            "vgp": vgp,
            "employedPops": emp if emp > 0 else 5000.0,
            "projectedRealGrowth": growth,
            "govtOwnedShares": govt_shares,
            "numberOfShares": tot_shares,
            "dividendEfficiency": div_eff,
            "dividendsPerShare": div_ps,
            "productionMethod": prod_method,
            "subsidised": False,
            "targetGrowth": 0.05,
            "illegal": comp_type == "Illegal Activities"
        })

ws_ss = wb['States Summary']
ss_rows = list(ws_ss.iter_rows(min_row=4, max_row=12, max_col=30, values_only=True))

vn_nations = []
for r in ss_rows:
    if r and len(r) > 1 and r[1] in state_names:
        st = r[1]
        pop = clean_val(r[9], 4000000.0)
        lit = clean_val(r[12], 0.65)
        urb = clean_val(r[14], 0.25)
        infra = clean_val(r[22], 0.35)
        vn_nations.append({
            "id": st.lower().replace(" ", "_"),
            "name": st,
            "planet": "Vietnam (1977 Regions)",
            "stockExchangeName": "Hanoi State Exchange",
            "consumerSentiment": 50.0,
            "kFactorScaler": 1500.0,
            "workingPops": pop if pop > 1000 else 4000000.0,
            "literacyRate": lit if 0 < lit <= 1 else 0.65,
            "urbanisationIndex": urb if 0 < urb <= 1 else 0.25,
            "infrastructureEfficiency": infra if 0 < infra <= 1 else 0.35,
            "rulingParty": "Communist Party (VCP)",
            "federalCooperativity": 0.3,
            "diplomaticExpenses": 0.0,
            "treasuryBalance": clean_val(r[11], 0.0),
            "specializedIndustry": "Agricultural Sector GDP" if "Delta" in st or "Highlands" in st else "Industry & Manufacturing GDP",
            "crimePolicy": "Beat Cops",
            "healthcarePolicy": "State-run Healthcare",
            "highImmigration": False,
            "companies": companies_by_state[st],
            "market": base_seed["nations"][0]["market"]
        })

vn_parties = [
    {
        "nation": "Vietnam",
        "name": "Communist Party (VCP)",
        "currentSupport": 0.85,
        "incomeTaxPreference": 0.35,
        "salesTaxPreference": 0.05,
        "corporateTaxRate": 0.50,
        "subsidyPreference": "Critical Goods",
        "nationalizationPreference": "Nationalize Critical Goods",
        "workerShareOfGdp": 0.65,
        "budgetAsPercentOfGdp": 0.25,
        "federalisedPolicy": "Unitary",
        "localCrimePolicy": "Beat Cops",
        "healthcarePolicy": "State-run Healthcare",
        "politicalAlignment": "Far left"
    },
    {
        "nation": "Vietnam",
        "name": "Reformist Faction (Doi Moi)",
        "currentSupport": 0.15,
        "incomeTaxPreference": 0.20,
        "salesTaxPreference": 0.10,
        "corporateTaxRate": 0.25,
        "subsidyPreference": "None",
        "nationalizationPreference": "No Nationalization",
        "workerShareOfGdp": 0.50,
        "budgetAsPercentOfGdp": 0.18,
        "federalisedPolicy": "Confederated",
        "localCrimePolicy": "Gentle Touch",
        "healthcarePolicy": "State-run Healthcare",
        "politicalAlignment": "Centre-Left"
    }
]

vn_seed = {
    "timeManager": {
        "intervalDaysPerHour": 90.0,
        "daysElapsed": 0.0,
        "quarterNumber": 1,
        "crisisStartDate": "1977-01-01",
        "crisisDate": "1977-01-01",
        "nextPaymentDate": "1977-04-01",
        "isLive": False
    },
    "politicalParties": vn_parties,
    "sectorsConfig": base_seed["sectorsConfig"],
    "techEras": base_seed["techEras"],
    "goodsCatalog": base_seed["goodsCatalog"],
    "nations": vn_nations,
    "stockAccounts": [
        {
            "id": "acc_vn_1",
            "holderName": "State Bank of Vietnam",
            "dividendPot": 0.0,
            "cashBalance": 800000000.0,
            "holdings": []
        }
    ]
}

with open('scripts/vietnam_seed.json', 'w') as f:
    json.dump(vn_seed, f, indent=2)

print(f"Vietnam seed extracted: {len(vn_nations)} states, {sum(len(n['companies']) for n in vn_nations)} companies.")
